#!/usr/bin/env python3
"""Build branded YASC technical whitepaper DOCX and PDF release artifacts.

Markdown sources are authoritative. DOCX/PDF are release artifacts. The build adds
Yofune branding supplied by the project owner, including cover mark, document
accent, headers/footers, website, and contact details.
"""
from __future__ import annotations
import shutil, subprocess
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[1]
DOCS = ROOT / "docs"
OUT = ROOT / "whitepaper"
REFERENCE = OUT / "reference.docx"
LOGO = ROOT / "assets" / "brand" / "yofune-mark-transparent.png"
VERSION = "1.0.0"
COMPANY = "Chengdu Yofune Ariake Technology Co., Ltd."
WEBSITE = "https://yofunesec.com/"
EMAIL = "contact@yofunesec.com"
BRAND = RGBColor(0x34, 0x46, 0x58)
MUTED = RGBColor(0x6E, 0x79, 0x84)

BUILDS = [
    (DOCS / "whitepaper.md", OUT / f"YASC-Technical-Whitepaper-v{VERSION}.docx", "en"),
    (DOCS / "whitepaper.zh-CN.md", OUT / f"YASC-Technical-Whitepaper-v{VERSION}.zh-CN.docx", "zh"),
]


def require(cmd: str) -> str:
    found = shutil.which(cmd)
    if not found:
        raise SystemExit(f"required executable not found: {cmd}")
    return found


def add_page_field(run):
    fld = OxmlElement('w:fldSimple')
    fld.set(qn('w:instr'), 'PAGE')
    r = OxmlElement('w:r')
    t = OxmlElement('w:t')
    t.text = '1'
    r.append(t); fld.append(r)
    run._r.addnext(fld)


def set_cell_margins(cell, top=0, start=0, bottom=0, end=0):
    tc = cell._tc; tcPr = tc.get_or_add_tcPr(); tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar'); tcPr.append(tcMar)
    for m,v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node=tcMar.find(qn(f'w:{m}'))
        if node is None: node=OxmlElement(f'w:{m}'); tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'),'dxa')


def style_document(doc: Document, lang: str):
    # Core metadata
    cp=doc.core_properties
    cp.title='Yofune Agent Security Checklist 2026 — Technical Whitepaper'
    cp.subject='Verifiable security baseline and verification methodology for enterprise AI agents'
    cp.author='Yofune Security Research'
    cp.keywords='AI agent security, YASC, verification, assurance, MCP, evidence'
    cp.comments=f'Published by {COMPANY}; {WEBSITE}; {EMAIL}'

    # Brand style accents while preserving the tested CJK-capable reference fonts.
    for name,size in [('Title',27),('Heading 1',18),('Heading 2',14),('Heading 3',12)]:
        try:
            s=doc.styles[name]; s.font.color.rgb=BRAND
            if name!='Title': s.font.bold=True
        except KeyError: pass
    try:
        doc.styles['Subtitle'].font.color.rgb=MUTED
    except KeyError: pass
    try:
        doc.styles['Hyperlink'].font.color.rgb=BRAND
    except KeyError: pass

    # Branded cover mark before the first content paragraph.
    if doc.paragraphs:
        first=doc.paragraphs[0]
        p=first.insert_paragraph_before()
        p.alignment=WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after=Pt(8)
        r=p.add_run(); r.add_picture(str(LOGO), width=Inches(1.35))

    # Give the final contact section a deliberate branded back-cover treatment.
    contact_titles = {"Contact Yofune", "联系 Yofune"}
    for idx, para in enumerate(doc.paragraphs):
        if para.text.strip() in contact_titles:
            para.paragraph_format.page_break_before = False
            para.alignment = WD_ALIGN_PARAGRAPH.CENTER
            para.paragraph_format.space_before = Pt(16)
            para.paragraph_format.space_after = Pt(14)
            mark = para.insert_paragraph_before()
            mark.alignment = WD_ALIGN_PARAGRAPH.CENTER
            mark.paragraph_format.page_break_before = True
            mark.paragraph_format.space_before = Pt(72)
            mark.paragraph_format.space_after = Pt(16)
            rr = mark.add_run(); rr.add_picture(str(LOGO), width=Inches(1.75))
            # Center the compact contact block and closing feedback on the back cover.
            refreshed = doc.paragraphs
            hidx = next(i for i, p2 in enumerate(refreshed) if p2._p is para._p)
            for following in refreshed[hidx+1:hidx+6]:
                try:
                    following.style = doc.styles["Body Text"]
                except KeyError:
                    pass
                following.alignment = WD_ALIGN_PARAGRAPH.CENTER
                following.paragraph_format.space_after = Pt(12)
            break

    for section in doc.sections:
        section.different_first_page_header_footer=True
        # Normal-page header
        header=section.header
        for p in header.paragraphs:
            p.clear()
        hp=header.paragraphs[0]
        hp.alignment=WD_ALIGN_PARAGRAPH.LEFT
        hp.paragraph_format.space_after=Pt(3)
        rr=hp.add_run(); rr.add_picture(str(LOGO), width=Inches(0.20))
        tr=hp.add_run(f"  YOFUNE SECURITY RESEARCH  |  YASC v{VERSION}")
        tr.font.size=Pt(8); tr.font.bold=True; tr.font.color.rgb=BRAND
        # First page header intentionally blank because the cover contains the mark.
        fheader=section.first_page_header
        for p in fheader.paragraphs: p.clear()

        # Footer with contact and page number.
        footer=section.footer
        for p in footer.paragraphs: p.clear()
        fp=footer.paragraphs[0]
        fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
        fp.paragraph_format.space_before=Pt(3)
        r=fp.add_run(f"{WEBSITE.replace('https://','')}  |  {EMAIL}  |  {COMPANY}  |  ")
        r.font.size=Pt(7.5); r.font.color.rgb=MUTED
        rp=fp.add_run(); rp.font.size=Pt(7.5); rp.font.color.rgb=MUTED; add_page_field(rp)
        # Cover footer: compact website/contact only.
        ff=section.first_page_footer
        for p in ff.paragraphs: p.clear()
        p0=ff.paragraphs[0]; p0.alignment=WD_ALIGN_PARAGRAPH.CENTER
        r0=p0.add_run(f"{WEBSITE.replace('https://','')}  •  {EMAIL}")
        r0.font.size=Pt(8); r0.font.color.rgb=MUTED


def main() -> None:
    pandoc = require("pandoc")
    office = shutil.which("soffice") or shutil.which("libreoffice")
    if not office:
        raise SystemExit("required executable not found: soffice/libreoffice")
    if not REFERENCE.exists():
        raise SystemExit(f"missing reference DOCX: {REFERENCE}")
    if not LOGO.exists():
        raise SystemExit(f"missing brand mark: {LOGO}")
    OUT.mkdir(exist_ok=True)

    for src, docx, lang in BUILDS:
        subprocess.run([
            pandoc, str(src), "--from=markdown",
            f"--reference-doc={REFERENCE}",
            f"--resource-path={DOCS}:{ROOT}",
            "-o", str(docx),
        ], check=True, cwd=ROOT)
        doc=Document(docx); style_document(doc,lang); doc.save(docx)
        subprocess.run([
            office, "--headless", "--convert-to", "pdf", "--outdir", str(OUT), str(docx)
        ], check=True, cwd=ROOT, stdout=subprocess.PIPE, stderr=subprocess.STDOUT, text=True)
        print(docx.relative_to(ROOT))
        print(docx.with_suffix('.pdf').relative_to(ROOT))

if __name__ == "__main__":
    main()
